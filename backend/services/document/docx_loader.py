"""
DOCX document loader using python-docx.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from backend.services.document.base_loader import BaseDocumentLoader
from backend.utils.logger import get_logger

logger = get_logger(__name__)


class DOCXLoader(BaseDocumentLoader):
    """
    Loader for Microsoft Word (.docx) documents.
    """

    def __init__(self, file_path: str | Path) -> None:
        super().__init__(file_path)

    async def load(self) -> str:
        """
        Extract text from a DOCX document.

        Returns:
            Extracted document text.

        Raises:
            Exception:
                If document loading fails.
        """
        self.validate()

        logger.info("Reading DOCX: %s", self.file_path)

        try:
            document = Document(self.file_path)

            content: list[str] = []

            # Paragraphs
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if text:
                    content.append(text)

            # Tables
            for table in document.tables:
                for row in table.rows:
                    row_data = []

                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        row_data.append(cell_text)

                    if any(row_data):
                        content.append(" | ".join(row_data))

            extracted_text = "\n".join(content).strip()

            logger.info(
                "Successfully extracted %d characters from '%s'.",
                len(extracted_text),
                self.filename,
            )

            return extracted_text

        except Exception:
            logger.exception(
                "Failed to extract text from DOCX '%s'.",
                self.filename,
            )
            raise

    def document_metadata(self) -> dict:
        """
        Return DOCX metadata.

        Returns:
            Metadata dictionary.
        """
        self.validate()

        try:
            document = Document(self.file_path)
            props = document.core_properties

            return {
                **super().metadata(),
                "title": props.title,
                "author": props.author,
                "subject": props.subject,
                "keywords": props.keywords,
                "category": props.category,
                "comments": props.comments,
                "created": props.created,
                "modified": props.modified,
                "last_modified_by": props.last_modified_by,
                "revision": props.revision,
            }

        except Exception:
            logger.exception(
                "Failed to read metadata from '%s'.",
                self.filename,
            )
            raise